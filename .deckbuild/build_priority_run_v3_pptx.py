#!/usr/bin/env python3
"""PI checkpoint deck — concise bookmark, no pipeline/code jargon."""
import importlib.util
import json
import sys
from pathlib import Path

from pptx.util import Inches

ROOT = Path(__file__).parent
DATA = json.loads((ROOT / "priority_run_v3_data.json").read_text(encoding="utf-8"))
OUT = ROOT.parent / "data" / "gene_sets" / DATA["run"] / "Priority_Gene_Sets_Overview.pptx"
TOTAL = 6

PI_BLURB = {
    "AD · DIOPT Bellenguez/Lambert": "Alzheimer's GWAS genes mapped to fly orthologs",
    "BPD · GWAS DIOPT Filter 2": "Bipolar disorder GWAS fly orthologs (Nature 2025)",
    "GABA gene list": "GABAergic genes from vGAT bulk-seq screen",
    "Slumber · Jordan VCM": "Curated VCM candidate genes (Slumber)",
    "Slumber · neuropeptides": "Neuropeptide & transmitter genes (Slumber)",
    "Slumber · RNAi hits": "RNAi screen hit candidates (Slumber)",
    "Tx-Omics mechanistic (nsyb/elav)": "Mechanistic screen candidates (Rosensweig–Shah 2026)",
}

BUCKET_LABEL = {
    "Bloomington >> RNAi_reagent >> Phenotype": "Bloomington · RNAi",
    "Vienna >> RNAi_reagent >> Phenotype": "Vienna · RNAi",
    "Bloomington >> AlleleOrInsertion >> Phenotype": "Bloomington · Allele",
    "Stock Center >> AlleleOrInsertion >> Phenotype": "Other stock centers · Allele",
    "Non Stock Center >> RNAi_reagent >> Phenotype": "Non–stock-center · RNAi",
    "Non Stock Center >> AlleleOrInsertion >> Phenotype": "Non–stock-center · Allele",
}

spec = importlib.util.spec_from_file_location("deck_helpers", ROOT / "build_pptx.py")
bp = importlib.util.module_from_spec(spec)
sys.modules["deck_helpers"] = bp
spec.loader.exec_module(bp)

W, H = bp.W, bp.H
INK, TEAL, TEALL, AMBER = bp.INK, bp.TEAL, bp.TEALL, bp.AMBER
WHITE, BG, TX, MUT, LINE, LIGHT, PALE = bp.WHITE, bp.BG, bp.TX, bp.MUT, bp.LINE, bp.LIGHT, bp.PALE

blank_slide = bp.blank_slide
fill_bg = bp.fill_bg
header = bp.header
footer = bp.footer
textbox = bp.textbox
rect = bp.rect
stat_card = bp.stat_card
add_table = bp.add_table


def bucket_rows(buckets):
    return [[BUCKET_LABEL.get(b[0], b[0]), b[1]] for b in buckets]


def compact_bucket_table(slide, left, top, width, buckets):
    add_table(
        slide, left, top, width,
        ["Bucket", "Stocks"],
        bucket_rows(buckets),
        [width - Inches(1.05), Inches(1.05)],
        row_size=9.5,
        header_size=10,
    )


def slide_title(prs):
    slide = blank_slide(prs)
    fill_bg(slide, INK)
    rect(slide, Inches(0), Inches(0), Inches(0.28), H, TEAL)
    rect(slide, Inches(0.28), Inches(0), Inches(0.08), H, AMBER)
    textbox(slide, Inches(1), Inches(1.6), Inches(11), Inches(0.4), "ALLADA LAB · REAGENT STOCKER", size=14, bold=True, color=TEALL)
    textbox(slide, Inches(0.95), Inches(2.1), Inches(11.6), Inches(0.9), "Priority Gene Sets", size=44, bold=True, color=WHITE, font="Georgia")
    textbox(slide, Inches(1), Inches(3.05), Inches(11), Inches(0.45), "Checkpoint · June 2026 priority run", size=18, color=LIGHT)
    items = [
        (str(len(DATA["gene_sets"])), "gene lists"),
        ("6", "stock buckets per list"),
        (f"{DATA['grand_total_stocks']:,}", "phenotypic stocks total"),
    ]
    x = Inches(1.2)
    for val, lab in items:
        textbox(slide, x, Inches(4.2), Inches(3.2), Inches(0.65), val, size=34, bold=True, color=AMBER, font="Georgia")
        textbox(slide, x, Inches(4.9), Inches(3.2), Inches(0.45), lab, size=13, color=LIGHT)
        x += Inches(3.5)


def slide_overview(prs):
    slide = blank_slide(prs)
    fill_bg(slide)
    header(slide, "At a glance", "Seven gene lists and phenotypic stock totals")
    rows = []
    for i, g in enumerate(DATA["gene_sets"], 1):
        rows.append([
            str(i),
            {"t": g["short"], "c": TEAL},
            str(g["n_genes"]),
            str(g["total"][1]),
            PI_BLURB.get(g["short"], g["provenance"]),
        ])
    add_table(
        slide, Inches(0.55), Inches(1.48), Inches(12.25),
        ["#", "Gene list", "Genes", "Stocks", "Source"],
        rows,
        [Inches(0.4), Inches(2.8), Inches(0.85), Inches(0.85), Inches(7.35)],
        row_size=10.5,
        header_size=11,
    )
    footer(slide, 2)


def slide_workflow(prs):
    slide = blank_slide(prs)
    fill_bg(slide)
    header(slide, "How lists become stock tables", "Three steps · each stock appears in one bucket only")

    steps = [
        ("1", "Gene lists", "GWAS orthologs, screen hits, and Slumber / Tx-Omics picks"),
        ("2", "Phenotypic stocks", "FlyBase reagents tied to a curated phenotype for each gene"),
        ("3", "Six buckets", "Sorted by stock center and reagent type (RNAi vs allele/insertion)"),
    ]
    x = Inches(0.55)
    cw = Inches(3.95)
    for num, title, body in steps:
        rect(slide, x, Inches(1.55), cw, Inches(1.35), WHITE, LINE)
        oval = slide.shapes.add_shape(bp.MSO_SHAPE.OVAL, x + Inches(0.2), Inches(1.72), Inches(0.48), Inches(0.48))
        oval.fill.solid()
        oval.fill.fore_color.rgb = TEAL
        oval.line.fill.background()
        textbox(slide, x + Inches(0.2), Inches(1.72), Inches(0.48), Inches(0.48), num, size=17, bold=True, color=WHITE, font="Georgia", align=bp.PP_ALIGN.CENTER)
        textbox(slide, x + Inches(0.82), Inches(1.74), cw - Inches(1.0), Inches(0.4), title, size=14, bold=True, color=INK)
        textbox(slide, x + Inches(0.22), Inches(2.2), cw - Inches(0.4), Inches(0.6), body, size=10.5, color=MUT)
        if num != "3":
            textbox(slide, x + cw - Inches(0.05), Inches(1.95), Inches(0.35), Inches(0.5), "→", size=22, bold=True, color=AMBER, align=bp.PP_ALIGN.CENTER)
        x += cw + Inches(0.22)

    textbox(slide, Inches(0.55), Inches(3.15), Inches(12), Inches(0.32), "The six priority buckets", size=13, bold=True, color=INK)
    buckets_plain = [
        "Bloomington · RNAi knockdown",
        "Vienna (VDRC) · RNAi knockdown",
        "Bloomington · classical alleles / insertions",
        "Other stock centers · alleles / insertions",
        "Lab / custom · RNAi reagents",
        "Lab / custom · alleles / insertions",
    ]
    bw, bh = Inches(3.95), Inches(0.72)
    for i, label in enumerate(buckets_plain):
        col, row = i % 3, i // 3
        bx = Inches(0.55) + col * (bw + Inches(0.18))
        by = Inches(3.55) + row * (bh + Inches(0.14))
        rect(slide, bx, by, Inches(0.12), bh, bp.BC[i] if i < len(bp.BC) else TEAL)
        rect(slide, bx + Inches(0.12), by, bw - Inches(0.12), bh, WHITE, LINE)
        textbox(slide, bx + Inches(0.28), by + Inches(0.14), bw - Inches(0.45), bh - Inches(0.2), label, size=10.5, color=TX)

    textbox(
        slide, Inches(0.55), Inches(6.35), Inches(12.2), Inches(0.55),
        "Only phenotype-backed stocks are kept. Literature on sleep, circadian rhythm, and locomotion is flagged when present.",
        size=10.5, color=MUT,
    )
    footer(slide, 3)


def slide_breakdowns(prs, n, title, gene_slice):
    slide = blank_slide(prs)
    fill_bg(slide)
    header(slide, "Stock breakdowns", title)
    sets = DATA["gene_sets"][gene_slice]
    ncols = len(sets)
    col_w = Inches(12.25 / ncols)
    x0 = Inches(0.55)
    for j, g in enumerate(sets):
        x = x0 + col_w * j
        rect(slide, x, Inches(1.45), col_w - Inches(0.12), Inches(5.55), WHITE, LINE)
        textbox(slide, x + Inches(0.14), Inches(1.58), col_w - Inches(0.35), Inches(0.55), g["short"], size=11.5, bold=True, color=TEAL)
        textbox(
            slide, x + Inches(0.14), Inches(2.05), col_w - Inches(0.35), Inches(0.45),
            f"{g['n_genes']} genes · {g['total'][1]} stocks", size=10, bold=True, color=INK,
        )
        textbox(slide, x + Inches(0.14), Inches(2.45), col_w - Inches(0.35), Inches(0.55), PI_BLURB.get(g["short"], ""), size=9, color=MUT)
        compact_bucket_table(slide, x + Inches(0.12), Inches(3.05), col_w - Inches(0.35), g["buckets"])
    footer(slide, n)


def friendly_total_label(combo_short: str) -> str:
    return (
        combo_short.replace("RNAi_reagent", "RNAi")
        .replace("Allele/Insertion", "Allele")
        .replace(" · Phenotype", "")
        .replace("Stock Center · Allele", "Other stock centers · Allele")
        .replace("Non-Stock-Center", "Non–stock-center")
        .replace("Vienna (VDRC)", "Vienna")
    )


def slide_totals(prs):
    slide = blank_slide(prs)
    fill_bg(slide)
    header(slide, "Combined totals", "All seven lists · phenotypic stocks by bucket")
    rows = [[friendly_total_label(t["combo_short"]), str(t["num_stocks"])] for t in DATA["totals_by_combination"]]
    add_table(
        slide, Inches(0.55), Inches(1.48), Inches(6.5),
        ["Bucket", "Stocks"],
        rows,
        [Inches(4.6), Inches(1.2)],
        row_size=11,
        header_size=11.5,
    )
    stat_card(slide, Inches(7.5), Inches(2.2), f"{DATA['grand_total_stocks']:,}", "total stocks (all lists)", AMBER)
    textbox(
        slide, Inches(7.5), Inches(3.5), Inches(5.3), Inches(1.2),
        "Full stock workbooks and the combined count summary are in the June 2026 priority run folder.",
        size=11, color=MUT,
    )
    footer(slide, TOTAL)


def main():
    bp.TOTAL = TOTAL
    prs = bp.new_prs()
    prs.core_properties.title = "Priority Gene Sets — Checkpoint"
    prs.core_properties.author = "Allada Lab"

    slide_title(prs)
    slide_overview(prs)
    slide_workflow(prs)
    slide_breakdowns(prs, 4, "Lists 1–4 · GWAS, GABA, and Slumber VCM", slice(0, 4))
    slide_breakdowns(prs, 5, "Lists 5–7 · Slumber and Tx-Omics", slice(4, 7))
    slide_totals(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    from sanitize_pptx import sanitize
    sanitize(OUT)
    print(f"WROTE {OUT}")


if __name__ == "__main__":
    main()
